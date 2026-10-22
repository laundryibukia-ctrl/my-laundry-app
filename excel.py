"""Excel export of reports (editable) and Excel -> PDF rendering."""

import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

RP = '"Rp "#,##0;-"Rp "#,##0'
HEAD_FILL = PatternFill("solid", fgColor="D0F8FF")
BOLD = Font(bold=True)


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], money_cols: set[int], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for idx in money_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell.number_format = RP
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def report_xlsx(summary, orders: list[dict], payments: list[dict], expenses: list[dict], labels: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    s = summary
    rows = [
        ["LAPORAN PENDAPATAN (nilai nota)", None], ["Jumlah pesanan", s.order_count], ["Penjualan kotor", s.gross_sales],
        ["Diskon", -s.discounts], ["Pendapatan bersih", s.revenue], [None, None],
        ["ARUS KAS (uang diterima/dikeluarkan)", None],
        *[[f"Kas masuk - {labels['method'].get(k, k)}", v] for k, v in s.cash_in_by_method.items()],
        ["Total kas masuk", s.cash_in],
        *[[f"Kas keluar - {labels['category'].get(k, k)}", -v] for k, v in s.cash_out_by_category.items()],
        ["Total kas keluar", -s.cash_out], ["Arus kas bersih", s.net_cash], [None, None],
        ["LABA-RUGI SEDERHANA", None], ["Pendapatan bersih", s.revenue], ["Beban bahan (stok terpakai)", -s.material_cost],
        ["Beban operasional", -s.operating_expenses], ["Beban gaji", -s.payroll_expense], ["Laba bersih", s.net_profit],
        [None, None], ["Piutang saat ini", s.receivables_now],
    ]
    ws = _sheet(wb, "Ringkasan", [f"Periode {s.date_from} s/d {s.date_to}", "Nilai"], rows, {2}, [44, 20])
    ws["B3"].number_format = "0"
    for (a, _b) in ws.iter_rows(min_row=2):
        if a.value and (str(a.value).isupper() or str(a.value).startswith(("Total", "Pendapatan bersih", "Laba", "Arus"))):
            a.font = BOLD
    _sheet(wb, "Harian", ["Tanggal", "Pendapatan", "Kas Masuk", "Kas Keluar"],
           [[d.date, d.revenue, d.cash_in, d.cash_out] for d in s.daily], {2, 3, 4}, [14, 18, 18, 18])
    _sheet(wb, "Pesanan", ["No Nota", "Tanggal", "Pelanggan", "No HP", "Status", "Total", "Dibayar", "Sisa"], orders, {6, 7, 8}, [18, 18, 26, 16, 14, 16, 16, 16])
    _sheet(wb, "Pembayaran", ["Tanggal", "No Nota", "Pelanggan", "Metode", "Jenis", "Nominal", "Status"], payments, {6}, [18, 18, 26, 14, 12, 16, 10])
    _sheet(wb, "Pengeluaran", ["Tanggal", "Kategori", "Keterangan", "Metode", "Nominal", "Status"], expenses, {5}, [12, 18, 36, 14, 16, 10])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Ya" if v else "Tidak"
    if isinstance(v, (int, float)):
        if "Rp" in (cell.number_format or ""):
            return ("-Rp " if v < 0 else "Rp ") + f"{abs(round(v)):,}".replace(",", ".")
        if isinstance(v, float) and not v.is_integer():
            return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"{int(v):,}".replace(",", ".") if abs(v) >= 10000 else str(int(v))
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return str(v)


def xlsx_to_pdf(data: bytes, title: str) -> bytes:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    buf = io.BytesIO()
    page = landscape(A4)
    margin = 12 * mm
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin, title=title)
    width = page[0] - 2 * margin
    st = getSampleStyleSheet()
    cell_style = st["BodyText"].clone("cell", fontSize=7.5, leading=9)
    story: list = []
    for i, ws in enumerate(wb.worksheets):
        rows = []
        for row in ws.iter_rows(max_row=min(ws.max_row, 3000)):
            vals = [_fmt(c) for c in row]
            if any(vals):
                rows.append(vals)
        if not rows:
            continue
        ncols = max(len(r) for r in rows)
        while ncols > 1 and all(len(r) < ncols or not r[ncols - 1] for r in rows):
            ncols -= 1
        rows = [(r + [""] * ncols)[:ncols] for r in rows]
        if i and story:
            story.append(PageBreak())
        story += [Paragraph(f"<b>{escape(title)}</b> — {escape(ws.title)}", st["Heading3"]), Spacer(1, 2 * mm)]
        widths = [max(8, min(45, max(len(r[c]) for r in rows))) for c in range(ncols)]
        scale = width / sum(widths)
        table = Table([[Paragraph(escape(v), cell_style) for v in r] for r in rows], colWidths=[w * scale for w in widths], repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D0F8FF")),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8E7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
    if not story:
        raise ValueError("File Excel kosong")
    doc.build(story)
    return buf.getvalue()
